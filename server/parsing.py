"""把 PDF / DOCX / PPTX / TXT 抽成按页（或按伪页）的纯文本。"""
from __future__ import annotations

import re
from pathlib import Path


def extract_pages(path: Path, ext: str) -> list[str]:
    ext = (ext or "").lower()
    if ext == ".pdf":
        return _pdf(path)
    if ext == ".docx":
        return _docx(path)
    if ext == ".pptx":
        return _pptx(path)
    if ext in (".txt", ".md", ".markdown", ".csv"):
        return [_read_text(path)]
    raise ValueError(f"暂不支持的文件类型：{ext or '未知'}")


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "ignore")


def _pdf(path: Path) -> list[str]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    return pages


def _docx(path: Path) -> list[str]:
    from docx import Document

    doc = Document(str(path))
    paras: list[str] = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                paras.append(" | ".join(cells))
    return _paginate(paras)


def _pptx(path: Path) -> list[str]:
    from pptx import Presentation

    prs = Presentation(str(path))
    pages: list[str] = []
    for slide in prs.slides:
        texts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text.strip()
                if text:
                    texts.append(text)
        pages.append("\n".join(texts))
    return pages


def _paginate(paras: list[str], size: int = 3000) -> list[str]:
    pages: list[str] = []
    buf = ""
    for p in paras:
        if buf and len(buf) + len(p) > size:
            pages.append(buf)
            buf = ""
        buf += p + "\n"
    if buf.strip():
        pages.append(buf)
    return pages


def clean(text: str) -> str:
    text = text.replace("\u00a0", " ").replace("\ufeff", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
